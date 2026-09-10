"""API tests for the review app.

Extraction is stubbed, so these run without an Anthropic key and assert the
review workflow: upload, list, fetch, edit (re-validation), delete, and export.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from voltscope.api.app import create_app
from voltscope.api.store import InMemoryBillStore

FIXTURES = Path(__file__).parent / "fixtures"
MESSY = json.loads((FIXTURES / "messy_bill.json").read_text(encoding="utf-8"))


def make_stub(returns=MESSY, error=None):
    """Return an extract_fn stub with the (data, filename) signature."""
    def _stub(data: bytes, filename: str):
        if error:
            return None, error
        return json.loads(json.dumps(returns)), None  # deep copy per call
    return _stub


@pytest.fixture
def client():
    app = create_app(store=InMemoryBillStore(), extract_fn=make_stub())
    return TestClient(app)


def _upload(client, name="bill1.pdf"):
    return client.post("/api/bills", files={"files": (name, b"%PDF-fake", "application/pdf")})


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_upload_processes_and_validates(client):
    r = _upload(client)
    assert r.status_code == 200
    rec = r.json()[0]
    assert rec["status"] == "processed"
    assert rec["filename"] == "bill1.pdf"
    codes = {f["code"] for f in rec["findings"]}
    # the messy fixture fires these
    assert "mpan_digit_count" in codes
    assert "subtotal_reconciliation" in codes


def test_list_and_get(client):
    rid = _upload(client).json()[0]["id"]
    assert any(b["id"] == rid for b in client.get("/api/bills").json())
    assert client.get(f"/api/bills/{rid}").json()["id"] == rid
    assert client.get("/api/bills/does-not-exist").status_code == 404


def test_edit_reruns_validation(client):
    rec = _upload(client).json()[0]
    rid = rec["id"]
    bill = rec["bill"]
    # Fix the MPAN to 13 digits -> the mpan finding should disappear.
    bill["supply_points"][0]["mpan"] = "1234567890123"
    updated = client.put(f"/api/bills/{rid}", json=bill).json()
    assert updated["status"] == "edited"
    codes = {f["code"] for f in updated["findings"]}
    assert "mpan_digit_count" not in codes


def test_edit_missing_record_404(client):
    assert client.put("/api/bills/nope", json={}).status_code == 404


def test_failed_extraction_records_failure():
    app = create_app(store=InMemoryBillStore(), extract_fn=make_stub(error="boom"))
    c = TestClient(app)
    rec = _upload(c).json()[0]
    assert rec["status"] == "failed"
    assert rec["error"] == "boom"
    assert rec["findings"] == []


def test_export_json(client):
    rid = _upload(client).json()[0]["id"]
    r = client.get(f"/api/bills/{rid}/export.json")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    assert json.loads(r.content)["id"] == rid


def test_export_xlsx_single_and_all(client):
    rid = _upload(client).json()[0]["id"]
    for url in (f"/api/bills/{rid}/export.xlsx", "/api/export.xlsx"):
        r = client.get(url)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert set(wb.sheetnames) == {"Bills", "Supply Points", "Charges", "Findings"}
        # Findings sheet has a header plus at least one finding row.
        assert wb["Findings"].max_row >= 2


def test_delete(client):
    rid = _upload(client).json()[0]["id"]
    assert client.delete(f"/api/bills/{rid}").status_code == 204
    assert client.get(f"/api/bills/{rid}").status_code == 404
