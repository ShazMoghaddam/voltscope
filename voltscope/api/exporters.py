"""Export bill records to JSON, Excel, and CSV.

The four flattenings (bills, supply points, charges, findings) are defined once
as ``(headers, rows)`` builders and reused by both the Excel writer and the CSV
writer, so the two formats never drift. There are no formulas, so no
recalculation step is needed.
"""

from __future__ import annotations

import csv
import io
from typing import Any, Callable, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..models import BillRecord

_HEADER_FONT = Font(name="Arial", bold=True)

Table = Tuple[List[str], List[List[Any]]]


def record_to_json(record: BillRecord) -> Dict[str, Any]:
    """Return the JSON-serialisable form of a single record."""
    return record.model_dump(mode="json", by_alias=True)


def _bill_field(record: BillRecord, name: str) -> Any:
    bill = record.bill
    return getattr(bill, name, None) if bill is not None else None


def _severity_counts(record: BillRecord) -> Dict[str, int]:
    counts = {"ERROR": 0, "WARNING": 0, "INFO": 0}
    for f in record.findings:
        counts[f.severity.value] = counts.get(f.severity.value, 0) + 1
    return counts


def _sp_dict(sp: Any) -> Dict[str, Any]:
    return sp.model_dump(by_alias=True) if hasattr(sp, "model_dump") else dict(sp)


# --- table builders (shared by Excel and CSV) ---
def bills_table(records: List[BillRecord]) -> Table:
    headers = [
        "Bill ID", "Filename", "Status", "Supplier", "Customer", "Account",
        "Invoice #", "Invoice Date", "Period Start", "Period End",
        "Contract End", "Currency", "Subtotal", "VAT", "Total",
        "Sites", "Findings", "Errors", "Warnings", "Info", "Error Msg",
    ]
    rows = []
    for r in records:
        sc = _severity_counts(r)
        sites = _bill_field(r, "supply_points") or []
        rows.append([
            r.id, r.filename, r.status.value,
            _bill_field(r, "supplier_name"), _bill_field(r, "customer_name"),
            _bill_field(r, "account_number"), _bill_field(r, "invoice_number"),
            _bill_field(r, "invoice_date"), _bill_field(r, "billing_period_start"),
            _bill_field(r, "billing_period_end"), _bill_field(r, "contract_end_date"),
            _bill_field(r, "currency"), _bill_field(r, "subtotal_gbp"),
            _bill_field(r, "vat_gbp"), _bill_field(r, "total_gbp"),
            len(sites), len(r.findings), sc["ERROR"], sc["WARNING"], sc["INFO"],
            r.error,
        ])
    return headers, rows


def supply_points_table(records: List[BillRecord]) -> Table:
    headers = [
        "Bill ID", "Filename", "Site #", "Fuel", "Supply Address", "MPAN",
        "MPRN", "Meter Serial", "Tariff", "Standing Charge (p/day)",
    ]
    rows = []
    for r in records:
        for i, sp in enumerate(_bill_field(r, "supply_points") or []):
            d = _sp_dict(sp)
            rows.append([
                r.id, r.filename, i, d.get("fuel_type"), d.get("supply_address"),
                d.get("mpan"), d.get("mprn"), d.get("meter_serial"),
                d.get("tariff_name"), d.get("standing_charge_p_per_day"),
            ])
    return headers, rows


def charges_table(records: List[BillRecord]) -> Table:
    headers = ["Bill ID", "Filename", "Site #", "Description", "Amount (GBP)"]
    rows = []
    for r in records:
        for i, sp in enumerate(_bill_field(r, "supply_points") or []):
            for c in _sp_dict(sp).get("charges") or []:
                rows.append([r.id, r.filename, i, c.get("description"), c.get("amount_gbp")])
    return headers, rows


def findings_table(records: List[BillRecord]) -> Table:
    headers = ["Bill ID", "Filename", "Severity", "Category", "Code", "Message",
               "Affected Field", "Recommendation"]
    rows = []
    for r in records:
        for f in r.findings:
            rows.append([
                r.id, r.filename, f.severity.value, f.category.value, f.code,
                f.message, f.affected_field, f.recommendation,
            ])
    return headers, rows


TABLES: Dict[str, Callable[[List[BillRecord]], Table]] = {
    "bills": bills_table,
    "supply_points": supply_points_table,
    "charges": charges_table,
    "findings": findings_table,
}


# --- Excel ---
def _write_sheet(ws: Worksheet, table: Table) -> None:
    headers, rows = table
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append(row)
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 60)


def records_to_workbook(records: List[BillRecord]) -> Workbook:
    """Flatten records into a styled four-sheet workbook."""
    wb = Workbook()
    first = wb.active
    first.title = "Bills"
    _write_sheet(first, bills_table(records))
    for name, builder in (("Supply Points", supply_points_table),
                          ("Charges", charges_table),
                          ("Findings", findings_table)):
        _write_sheet(wb.create_sheet(name), builder(records))
    return wb


def records_to_xlsx_bytes(records: List[BillRecord]) -> bytes:
    """Serialise a workbook of records to xlsx bytes."""
    wb = records_to_workbook(records)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- CSV ---
def records_to_csv(records: List[BillRecord], sheet: str = "bills") -> str:
    """Return one of the flattenings as CSV text.

    ``sheet`` is one of ``bills``, ``supply_points``, ``charges``, ``findings``.
    """
    builder = TABLES.get(sheet)
    if builder is None:
        raise KeyError(sheet)
    headers, rows = builder(records)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue()
